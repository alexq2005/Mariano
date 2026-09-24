#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extraer.py - convierte el XLSX del proveedor en los datos del catalogo web.

Lee el Excel (productos + fotos embebidas), saca miniaturas y escribe
datos/proveedor.json (PRIVADO: incluye el costo en dolares) y las fotos en
public/img/. Despues hay que correr:

    npm run catalogo

que calcula los precios y genera public/data/catalogo.json, que es lo que
carga la pagina. El costo NUNCA se publica.

Las fotos van dentro del XLSX como dibujos anclados a una celda. openpyxl
no las expone, asi que el anclaje se lee del XML: xdr:from/xdr:row da la
fila, y r:embed apunta al archivo en xl/media/.

    python extraer.py <archivo.xlsx> [salida/]     (salida por defecto: la raiz del proyecto)
"""

import os
import re
import sys
import json
import shutil
import zipfile
import warnings

warnings.filterwarnings("ignore")

try:
    import openpyxl
    from PIL import Image
except ImportError as e:
    raise SystemExit(f"Falta una dependencia: {e}\n  pip install openpyxl Pillow")

HOJA = "Catálogo"
FILA_1 = 5          # primera fila con producto (1-indexada)
ANCHO_MINI = 340    # px; suficiente en celular y en desktop a 2x


# Los rubros salen del nombre del producto. Gana la primera que coincide,
# y el ORDEN IMPORTA: lo que dice que ES el producto (mascarilla, crema,
# toallita) tiene que evaluarse antes que lo que solo lo describe. "Crema
# iluminadora" es una crema, no maquillaje; con Rostro antes que Cuidado,
# "iluminador" se la llevaba a Rostro.
#
# Los \b tambien importan: sin limite de palabra, "set" matcheaba adentro de
# "MAKEUP SETTING SPRAY" y el spray fijador caia en Accesorios.
RUBROS = [
    ("labios",      r"labial|labios|gloss|brillo de labios|lip|bálsamo|balsamo"),
    ("ojos",        r"ceja|pestañ|delineador|sombra|eyeliner|máscara de pest|rímel|rimel"
                    r"|lente|gel fijador"),
    ("cuidado",     r"mascarilla|crema|sérum|serum|exfoliante|limpiador|aceite|protector"
                    r"|hidratante|ampolla|tónico|tonico|toallita|desmaquill|parche|antifaz"),
    ("rostro",      r"\bbase\b|corrector|polvo|rubor|contorno|iluminador|bb cream|primer"
                    r"|fijador de maquillaje|setting spray|glitter|multiuso"),
    ("uñas",        r"uña|esmalte|manicur"),
    ("cabello",     r"cabello|\bpelo\b|shampoo|champú|acondicionador|\bcera\b|canas"
                    r"|tinte para el cabello"),
    ("accesorios",  r"brocha|pincel|esponja|espejo|neceser|\bset\b|organizador|pinza"),
]


def rubro_de(nombre, descripcion):
    texto = f"{nombre} {descripcion}".lower()
    for rubro, patron in RUBROS:
        if re.search(patron, texto):
            return rubro
    return "otros"


def limpio(codigo):
    """ZMA-1310/07 -> ZMA-1310-07, apto como nombre de archivo y como URL."""
    return re.sub(r"[^A-Za-z0-9_-]", "-", str(codigo)).strip("-") or "sin-codigo"


# Excel convierte los codigos de barra largos a notacion cientifica
# ("7E+12"); como descripcion no le dicen nada a la clienta.
NOTACION_CIENTIFICA = re.compile(r"^\d+(\.\d+)?E\+\d+$", re.I)


def mapa_imagenes(ruta_xlsx):
    """{fila 0-indexada: 'xl/media/imageN.png'} leido del XML de dibujos."""
    with zipfile.ZipFile(ruta_xlsx) as z:
        try:
            dib = z.read("xl/drawings/drawing1.xml").decode("utf-8")
            rel = z.read("xl/drawings/_rels/drawing1.xml.rels").decode("utf-8")
        except KeyError:
            return {}

    rels = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rel))
    mapa = {}
    anclas = re.findall(
        r"<xdr:(?:oneCellAnchor|twoCellAnchor).*?</xdr:(?:oneCellAnchor|twoCellAnchor)>",
        dib, re.S,
    )
    for a in anclas:
        fila = re.search(r"<xdr:from>.*?<xdr:row>(\d+)</xdr:row>", a, re.S)
        emb = re.search(r'r:embed="([^"]+)"', a)
        if not (fila and emb and emb.group(1) in rels):
            continue
        destino = rels[emb.group(1)].replace("../", "xl/")
        mapa[int(fila.group(1))] = destino
    return mapa


def main():
    if len(sys.argv) < 2:
        raise SystemExit(f"Uso: python {os.path.basename(sys.argv[0])} <archivo.xlsx> [salida/]")

    xlsx = sys.argv[1]
    salida = sys.argv[2] if len(sys.argv) > 2 else "."
    # Las fotos son publicas; los datos con el costo, no: van a datos/,
    # que esta fuera del repositorio.
    carpeta_img = os.path.join(salida, "public", "img")
    carpeta_datos = os.path.join(salida, "datos")
    os.makedirs(carpeta_img, exist_ok=True)
    os.makedirs(carpeta_datos, exist_ok=True)

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[HOJA] if HOJA in wb.sheetnames else wb[wb.sheetnames[0]]
    imgs = mapa_imagenes(xlsx)
    print(f"Fotos ancladas en el Excel: {len(imgs)}")

    productos, sin_foto, sin_precio, vistos = [], 0, 0, set()

    with zipfile.ZipFile(xlsx) as z:
        for i, fila in enumerate(ws.iter_rows(min_row=FILA_1, values_only=True)):
            if fila[0] is None:
                continue
            fila_0 = FILA_1 - 1 + i          # la misma fila, 0-indexada

            codigo = str(fila[2] or "").strip()
            nombre = str(fila[3] or "").strip()
            if not nombre:
                continue

            desc = str(fila[4] or "").strip()
            if desc.lower() == nombre.lower() or NOTACION_CIENTIFICA.match(desc):
                desc = ""                      # nombre repetido o codigo de barras roto

            bulto = int(fila[5]) if isinstance(fila[5], (int, float)) else None
            costo = float(fila[6]) if isinstance(fila[6], (int, float)) else None
            if costo is None:
                sin_precio += 1
                continue                       # sin costo no se puede calcular precio

            slug = limpio(codigo)
            if slug in vistos:                 # hay codigos repetidos en el Excel
                slug = f"{slug}-{fila[0]}"
            vistos.add(slug)

            archivo = None
            origen = imgs.get(fila_0)
            if origen:
                try:
                    with z.open(origen) as fh:
                        im = Image.open(fh)
                        im.load()
                    if im.mode not in ("RGB", "L"):
                        fondo = Image.new("RGB", im.size, (255, 255, 255))
                        fondo.paste(im, mask=im.split()[-1] if im.mode in ("RGBA", "LA") else None)
                        im = fondo
                    if im.width > ANCHO_MINI:
                        alto = round(im.height * ANCHO_MINI / im.width)
                        im = im.resize((ANCHO_MINI, alto), Image.LANCZOS)
                    archivo = f"{slug}.jpg"
                    im.convert("RGB").save(
                        os.path.join(carpeta_img, archivo),
                        "JPEG", quality=78, optimize=True, progressive=True,
                    )
                except Exception as err:
                    print(f"  ! {codigo}: no se pudo procesar la foto ({err})")
                    archivo = None
            if archivo is None:
                sin_foto += 1

            productos.append({
                # id unico aunque el Excel repita el codigo: es la clave del
                # carrito y de la URL /product/<id>. "cod" queda tal cual,
                # porque es el codigo que entiende el proveedor.
                "id": slug,
                "cod": codigo,
                "nom": nombre,
                "desc": desc,
                "rubro": rubro_de(nombre, desc),
                "bulto": bulto,
                "costo": round(costo, 4),      # USD, NUNCA se muestra al cliente
                "img": archivo,
            })

    datos = {
        "generado_de": os.path.basename(xlsx),
        "productos": productos,
    }
    js = os.path.join(carpeta_datos, "proveedor.json")
    with open(js, "w", encoding="utf-8") as fh:
        json.dump(datos, fh, ensure_ascii=False, indent=2)
        fh.write("\n")

    peso = sum(
        os.path.getsize(os.path.join(carpeta_img, f)) for f in os.listdir(carpeta_img)
    ) / 1024 / 1024

    print(f"\nProductos ............. {len(productos)}")
    print(f"Sin precio (salteados)  {sin_precio}")
    print(f"Sin foto .............. {sin_foto}")
    print(f"Fotos ................. {len(os.listdir(carpeta_img))}  ({peso:.1f} MB)")
    print(f"Datos ................. {js}  ({os.path.getsize(js)/1024:.0f} KB)")
    rubros = {}
    for p in productos:
        rubros[p["rubro"]] = rubros.get(p["rubro"], 0) + 1
    print("\nPor rubro:")
    for r, n in sorted(rubros.items(), key=lambda x: -x[1]):
        print(f"  {r:<12s} {n:>4d}")

    print("\nFalta un paso: npm run catalogo  (calcula los precios y genera")
    print("public/data/catalogo.json, que es lo que carga la pagina).")


if __name__ == "__main__":
    main()
